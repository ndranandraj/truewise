"""pipeline/build_fvtge_source.py: ED's FVT/GE reporting spreadsheets, read row for row or refused.

Two kinds of test, and only the first is CI coverage:
  * Synthetic workbooks in both of ED's layouts (August's seven components with percentages only,
    October's ten with counts). Each is registered with its own SHA-256, so the checksum guard runs
    and every corruption reaches the rule it targets rather than being stopped earlier. These run
    everywhere, CI included.
  * The real releases, when ED's files are present locally (data/ is never committed, so these are
    SKIPPED in CI): August must reproduce published/fvtge_reporting.parquet value for value, and
    October must keep every row and give ED's counts. Evidence from these runs is recorded with the
    revision it ran on; it is not CI evidence.
"""

from __future__ import annotations

import dataclasses
from collections import Counter

import duckdb
import openpyxl
import pytest

from pipeline import build_fvtge_source as S

ROOT = S.ROOT


def _row(rel, op, name="A College", state="CA", control="Public", prior=None, current=None):
    comps = dict(zip(S.PRIOR, prior or ["Submitted"] * 7, strict=True))
    if rel is S.OCT:
        comps |= dict(zip(S.CURRENT, current or ["Submitted"] * 3, strict=True))
    out = {"opeid6": op, "instnm": name, "stabbr": state, "control": control, **comps}
    for num, (group, flag, yes, no) in rel.counts.items():
        n = sum(comps[c] == "Not Submitted" for c in group)
        out[num], out[flag] = n, (yes if n else no)
    return out


def _book(path, rel, rows, *, header=None, raw_rows=None, freq_edit=None, phrase=None):
    """A workbook in `rel`'s layout. Its Frequencies sheet is computed from `rows` (counts for October,
    percentages to two decimals for August, as ED gave them) unless edited."""
    cols = header or list(rel.columns)
    wb = openpyxl.Workbook()
    wb.active.title = "Overview"
    text = f"These data were {phrase or rel.compile_phrase}."
    wb["Overview"].append([text if rel.compile_sheet == "Overview" else "UPDATE: a test file."])
    if rel.compile_sheet == "Technical":
        wb.create_sheet("Technical").append([text])
    data = wb.create_sheet("Data")
    data.append(cols)
    for r in raw_rows if raw_rows is not None else [[d.get(c) for c in cols] for d in rows]:
        data.append(r)
    freq = wb.create_sheet("Frequencies")
    counted = rel is S.OCT
    freq.append(["Variable", "Outcome", *(["Count (N)"] if counted else []), "Frequency (%)"])
    freq.append(["opeid6", "Six-digit OPEID code", *(["N/A"] if counted else []), "N/A"])
    for var in [c for c in cols if c not in ("opeid6", "instnm", "stabbr")]:
        counts = Counter(str(d[var]) for d in rows if var in d)
        for i, (outcome, n) in enumerate(sorted(counts.items())):
            pct = round(100 * n / len(rows), 1 if counted else 2)
            freq.append([var if i == 0 else None, outcome, *([n] if counted else []), pct])
    if freq_edit:
        freq_edit(freq)
    wb.save(path)
    # Registered as itself: the checksum guard passes, so later rules are what is tested.
    return dataclasses.replace(rel, path=path, sha256=S.sha256(path))


def _oct_rows():
    return [
        _row(S.OCT, "001001"),
        _row(
            S.OCT,
            "001002",
            control="Private For-Profit",
            prior=["Not Submitted"] * 7,
            current=["Not Submitted"] * 3,
        ),
        _row(
            S.OCT,
            "001003",
            prior=["Not Required*"] + ["Submitted"] * 6,
            current=["Update: No Longer Required"] * 3,
        ),
    ]


def _aug_rows():
    return [
        _row(S.AUG, "002001"),
        _row(
            S.AUG,
            "002002",
            control="Foreign",
            prior=["Not Submitted"] * 2 + ["Not Required"] + ["Submitted"] * 4,
        ),
    ]


# ---- a valid file: every row and value kept ---------------------------------------------------


def test_october_layout_keeps_every_row_and_value_exactly(tmp_path):
    rows = _oct_rows()
    rel = _book(tmp_path / "ok.xlsx", S.OCT, rows)
    body = S.read_rows(rel)
    assert body == [{c: d[c] for c in S.OCT.columns} for d in rows]
    # Statuses verbatim, asterisk and all; leading zeros kept on every OPEID.
    assert (
        body[2]["total2223"] == "Not Required*"
        and body[2]["program2526"] == "Update: No Longer Required"
    )
    assert [d["opeid6"] for d in body] == ["001001", "001002", "001003"]
    out = tmp_path / "out.parquet"
    S.write(rel, body, out)
    got = duckdb.sql(f"SELECT * FROM '{out}' ORDER BY opeid6").fetchall()
    assert [r[0] for r in got] == ["001001", "001002", "001003"]
    cols = [c[0] for c in duckdb.sql(f"DESCRIBE SELECT * FROM '{out}'").fetchall()]
    assert cols == [*S.OCT.columns, "compiled", "published", "source_url", "source_sha256"]


def test_august_layout_has_no_2026_fields_rather_than_zeros(tmp_path):
    rows = _aug_rows()
    rel = _book(tmp_path / "aug.xlsx", S.AUG, rows)
    body = S.read_rows(rel)
    assert body == [{c: d[c] for c in S.AUG.columns} for d in rows]
    out = tmp_path / "aug.parquet"
    S.write(rel, body, out)
    cols = [c[0] for c in duckdb.sql(f"DESCRIBE SELECT * FROM '{out}'").fetchall()]
    assert not any(c in cols for c in (*S.CURRENT, "num_miss_current", "yes_missing_current"))
    assert cols == [*S.AUG.columns, "compiled", "source_url"]


def test_a_fully_empty_row_is_skipped_but_a_row_of_zeros_is_not(tmp_path):
    rows = _oct_rows()
    raw = [[d.get(c) for c in S.OCT.columns] for d in rows]
    raw.insert(1, [None] * len(S.OCT.columns))
    assert len(S.read_rows(_book(tmp_path / "gap.xlsx", S.OCT, rows, raw_rows=raw))) == 3
    raw[1] = [0] * len(S.OCT.columns)
    with pytest.raises(SystemExit, match=r"OPEID 0, not six digits"):
        S.read_rows(_book(tmp_path / "zeros.xlsx", S.OCT, rows, raw_rows=raw))


# ---- each rule refuses, reached directly -------------------------------------------------------


@pytest.mark.parametrize(
    ("change", "message"),
    [
        (lambda r: r[0].update(total2223="Submitted (late)"), "unknown submission status"),
        (lambda r: r[0].update(program2526=None), "unknown submission status"),
        (lambda r: r[1].update(opeid6="001001"), "OPEID 001001 appears twice"),
        (lambda r: r[0].update(opeid6="1001"), "not six digits"),
        (lambda r: r[0].update(opeid6="00100A"), "not six digits"),
        (lambda r: r[0].update(opeid6=1001), "not six digits"),
        (lambda r: r[0].update(instnm=None), "missing name or state"),
        (lambda r: r[0].update(control="Tribal"), "unknown control"),
        (lambda r: r[1].update(num_miss_prior=6), "num_miss_prior is 6"),
        # Excel stores every number as a float and openpyxl returns whole ones as int, so a whole
        # float cannot reach the parser; a fractional count can.
        (lambda r: r[1].update(num_miss_prior=6.5), "num_miss_prior is 6.5"),
        (lambda r: r[0].update(num_miss_current=False), "num_miss_current is False"),
        (lambda r: r[1].update(num_miss_total="10"), "num_miss_total is '10'"),
        (
            lambda r: r[0].update(yes_missing_current="HAS MISSING FILE(S) FROM CURRENT CYCLE"),
            "yes_missing_current is",
        ),
    ],
)
def test_a_bad_row_is_refused_by_its_own_rule(tmp_path, change, message):
    rows = _oct_rows()
    change(rows)
    rel = _book(tmp_path / "bad.xlsx", S.OCT, rows)
    with pytest.raises(SystemExit, match=message):
        S.read_rows(rel)


def test_a_row_with_values_but_no_opeid_is_refused_not_dropped(tmp_path):
    rows = _oct_rows()
    raw = [[d.get(c) for c in S.OCT.columns] for d in rows]
    raw[2][0] = None
    with pytest.raises(SystemExit, match="not six digits"):
        S.read_rows(_book(tmp_path / "noid.xlsx", S.OCT, rows, raw_rows=raw))


def test_a_changed_header_is_refused(tmp_path):
    cols = list(S.OCT.columns)
    cols[-1], cols[-2] = cols[-2], cols[-1]  # ED swapped two columns between August and September
    with pytest.raises(SystemExit, match="changed the Data sheet's columns"):
        S.read_rows(_book(tmp_path / "hdr.xlsx", S.OCT, _oct_rows(), header=cols))


def test_a_different_compile_date_is_refused(tmp_path):
    rel = _book(tmp_path / "date.xlsx", S.OCT, _oct_rows(), phrase="compiled on November 2nd, 2026")
    with pytest.raises(SystemExit, match="does not say"):
        S.read_rows(rel)


def test_the_real_registered_checksum_is_enforced(tmp_path):
    rel = _book(tmp_path / "other.xlsx", S.OCT, _oct_rows())
    with pytest.raises(SystemExit, match="SHA-256 differs"):
        S.read_rows(dataclasses.replace(rel, sha256=S.OCT.sha256))


def _edit_first(match):
    def edit(ws):
        for row in ws.iter_rows(min_row=3):
            if match(row):
                return row
        raise AssertionError("fixture row not found")

    return edit


def test_frequencies_counts_percentages_and_listing_are_each_checked(tmp_path):
    # October: a wrong count.
    def wrong_count(ws):
        _edit_first(lambda r: r[1].value == "Not Submitted" and isinstance(r[2].value, int))(ws)[
            2
        ].value += 1

    with pytest.raises(SystemExit, match="Frequencies sheet gives .* the rows give"):
        S.read_rows(_book(tmp_path / "f1.xlsx", S.OCT, _oct_rows(), freq_edit=wrong_count))

    # October: a value in the data that the sheet does not list (a non-first outcome, so the
    # variable's own name stays on the sheet).
    def drop_value(ws):
        for c in _edit_first(
            lambda r: r[1].value == "Update: No Longer Required" and r[0].value is None
        )(ws):
            c.value = None

    with pytest.raises(SystemExit, match="does not list"):
        S.read_rows(_book(tmp_path / "f2.xlsx", S.OCT, _oct_rows(), freq_edit=drop_value))

    # August: no counts, so the percentage is checked to ED's precision.
    def wrong_pct(ws):
        _edit_first(lambda r: r[1].value == "Not Submitted")(ws)[2].value = 12.34

    with pytest.raises(SystemExit, match=r"gives 12.34%"):
        S.read_rows(_book(tmp_path / "f3.xlsx", S.AUG, _aug_rows(), freq_edit=wrong_pct))


# ---- output safety -----------------------------------------------------------------------------


def test_a_refused_file_leaves_the_existing_output_untouched(tmp_path):
    out = tmp_path / "out.parquet"
    good = _book(tmp_path / "good.xlsx", S.OCT, _oct_rows())
    S.write(good, S.read_rows(good), out)
    before = out.read_bytes()
    rows = _oct_rows()
    rows[0]["total2223"] = "Pending"
    bad = _book(tmp_path / "bad.xlsx", S.OCT, rows)
    with pytest.raises(SystemExit):
        S.write(bad, S.read_rows(bad), out)
    assert out.read_bytes() == before


def test_a_write_that_does_not_read_back_identically_changes_nothing(tmp_path, monkeypatch):
    out = tmp_path / "out.parquet"
    rel = _book(tmp_path / "good.xlsx", S.OCT, _oct_rows())
    body = S.read_rows(rel)
    S.write(rel, body, out)
    before = out.read_bytes()
    real = S.duckdb.sql

    def lossy(q, *a, **k):
        r = real(q, *a, **k)
        if q.startswith("SELECT * FROM read_parquet"):
            r = real("SELECT * FROM (VALUES (1)) WHERE false")
        return r

    monkeypatch.setattr(S.duckdb, "sql", lossy)
    with pytest.raises(SystemExit, match="did not read back identically"):
        S.write(rel, body, out)
    assert out.read_bytes() == before
    assert not [p for p in tmp_path.iterdir() if ".tmp-" in p.name], (
        "a temporary file was left behind"
    )


def test_october_without_an_output_path_fails_before_reading_or_writing(tmp_path, monkeypatch):
    monkeypatch.setattr(
        S, "read_rows", lambda *a, **k: pytest.fail("read before the output was decided")
    )
    with pytest.raises(SystemExit, match="pass --out"):
        S.main(["--release", "2026-10-05"])
    assert S.AUG.default_out == ROOT / "published" / "fvtge_reporting.parquet"
    assert S.OCT.default_out is None


# ---- the real releases: local evidence only (skipped in CI) ------------------------------------


@pytest.mark.skipif(not S.AUG.path.exists(), reason="ED's August file is not in data/raw here")
def test_august_reproduces_the_published_file_exactly(tmp_path):
    out = tmp_path / "aug.parquet"
    S.write(S.AUG, S.read_rows(S.AUG), out)
    pub = ROOT / "published" / "fvtge_reporting.parquet"
    a = duckdb.sql(f"SELECT * FROM '{pub}' ORDER BY opeid6").fetchall()
    b = duckdb.sql(f"SELECT * FROM '{out}' ORDER BY opeid6").fetchall()
    types = lambda p: duckdb.sql(f"DESCRIBE SELECT * FROM '{p}'").fetchall()  # noqa: E731
    assert len(a) == 4635 and a == b and types(pub) == types(out)


@pytest.mark.skipif(not S.OCT.path.exists(), reason="ED's October file is not in data/raw here")
def test_october_keeps_every_row_and_agrees_with_ed():
    body = S.read_rows(S.OCT)
    rows = list(
        openpyxl.load_workbook(S.OCT.path, read_only=True)["Data"].iter_rows(values_only=True)
    )
    raw = {r[0]: dict(zip(rows[0], r, strict=True)) for r in rows[1:] if r and r[0]}
    assert len(body) == len(raw) == 4674
    assert all(raw[d["opeid6"]] == d for d in body)
    count = lambda f: sum(1 for d in body if f(d))  # noqa: E731
    assert count(lambda d: d["num_miss_prior"] > 0) == 1396
    assert count(lambda d: d["num_miss_prior"] == 7) == 478
    assert count(lambda d: d["num_miss_current"] > 0) == 1007
    assert count(lambda d: d["num_miss_current"] == 3) == 633
