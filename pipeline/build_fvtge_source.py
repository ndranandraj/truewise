"""Turn ED's FVT/GE reporting-status spreadsheet into a small committed source file.

ED attached "List of Institutions That Previously Submitted FVT/GE Data" to electronic announcement
GENERAL-26-49 (11 August 2026). It lists, for every open institution with at least one program of 30
or more students, which of the seven required FVT/GE file components it had submitted as of the date
ED compiled it (6 August 2026).

This step runs where the spreadsheet has been downloaded (it is not fetched by the build) and writes
published/fvtge_reporting.parquet, which is committed, checksummed in published/SHA256SUMS.txt, and
is what the deploy builds the finding from. It refuses to run if ED's columns or status vocabulary
change, rather than mapping a new layout onto old meanings.

Usage (from repo root, with the spreadsheet in data/raw/):
    python -m pipeline.build_fvtge_source
"""

from __future__ import annotations

import duckdb
import openpyxl

from pipeline.config import RAW_DIR, ROOT

SOURCE_XLSX = RAW_DIR / "FVTGEDataReportingFinal.xlsx"
OUT = ROOT / "published" / "fvtge_reporting.parquet"
SOURCE_URL = "https://fsapartners.ed.gov/sites/default/files/2026-08/FVTGEDataReportingFinal.xlsx"
COMPILED = "2026-08-06"  # from the spreadsheet's own Overview sheet: "compiled on August 6th, 2026"

COMPONENTS = (
    "total2223",
    "program2324",
    "annual2324",
    "total2324",
    "program2425",
    "annual2425",
    "total2425",
)
COLUMNS = ("opeid6", "instnm", "stabbr", "control", *COMPONENTS, "yes_missing", "num_miss")
STATUSES = {"Submitted", "Not Submitted", "Not Required"}


def read_rows() -> list[tuple]:
    wb = openpyxl.load_workbook(SOURCE_XLSX, read_only=True)
    overview = " ".join(str(c) for r in wb["Overview"].iter_rows(values_only=True) for c in r if c)
    if "compiled on August 6th, 2026" not in overview:
        raise SystemExit(
            "The spreadsheet's compile date is not 6 August 2026. This is a different release: update "
            "COMPILED and re-read ED's caveats before publishing anything from it."
        )
    rows = list(wb["Data"].iter_rows(values_only=True))
    if tuple(rows[0]) != COLUMNS:
        raise SystemExit(f"ED changed the Data sheet's columns: {rows[0]}")
    body = [tuple(r) for r in rows[1:] if r and r[0]]
    for r in body:
        bad = {r[i] for i in range(4, 11)} - STATUSES
        if bad:
            raise SystemExit(f"Unknown submission status {bad} for OPEID {r[0]}")
        missing = sum(r[i] == "Not Submitted" for i in range(4, 11))
        if missing != r[12]:
            raise SystemExit(f"num_miss disagrees with the component columns for OPEID {r[0]}")
    return body


def main() -> None:
    body = read_rows()
    con = duckdb.connect()
    cols = ", ".join(f"{c} VARCHAR" if c != "num_miss" else "num_miss INTEGER" for c in COLUMNS)
    con.execute(f"CREATE TABLE ed ({cols})")
    con.executemany(f"INSERT INTO ed VALUES ({', '.join('?' * len(COLUMNS))})", body)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    con.execute(
        f"COPY (SELECT *, '{COMPILED}' AS compiled, '{SOURCE_URL}' AS source_url FROM ed "
        f"ORDER BY opeid6) TO '{OUT}' (FORMAT PARQUET)"
    )
    n_miss = sum(r[11] == "HAS MISSING FILES" for r in body)
    print(
        f"FVT/GE status: {len(body):,} institutions, {n_miss:,} with at least one component not submitted"
    )
    print(f"wrote -> {OUT}")


if __name__ == "__main__":
    main()
